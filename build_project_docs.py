from __future__ import annotations

from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, Preformatted, SimpleDocTemplate, Spacer

ROOT = Path(r"c:/Users/gorla/GitHub Projects/RAG_Assistant")
MD_IN = ROOT / "Project_Documentation.md"
DOCX_OUT = ROOT / "Project_Documentation.docx"
PDF_OUT = ROOT / "Project_Documentation.pdf"

text = MD_IN.read_text(encoding="utf-8")
lines = text.splitlines()

# DOCX output

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Project Documentation").italic = True

for line in lines:
    if line.startswith("```"):
        continue
    if line.startswith("# "):
        doc.add_heading(line[2:], level=1)
    elif line.startswith("## "):
        doc.add_heading(line[3:], level=2)
    elif line.startswith("### "):
        doc.add_heading(line[4:], level=3)
    elif not line.strip():
        doc.add_paragraph("")
    else:
        p = doc.add_paragraph()
        p.add_run(line)

doc.save(DOCX_OUT)

# PDF output
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name="TitleCustom", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, leading=24, alignment=TA_CENTER, textColor=colors.HexColor("#0F4C81")))
styles.add(ParagraphStyle(name="H1Custom", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=20, spaceAfter=8, textColor=colors.HexColor("#0F4C81")))
styles.add(ParagraphStyle(name="H2Custom", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=16, spaceAfter=6))
styles.add(ParagraphStyle(name="BodyCustom", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=12))
styles.add(ParagraphStyle(name="CodeCustom", parent=styles["Code"], fontName="Courier", fontSize=7.5, leading=9))

story = [Paragraph("Project Documentation: Production RAG Knowledge Assistant", styles["TitleCustom"]), Spacer(1, 10)]
code_mode = False
for line in lines:
    if line.startswith("```"):
        code_mode = not code_mode
        continue
    if code_mode:
        story.append(Preformatted(line, styles["CodeCustom"]))
        continue
    if line.startswith("# "):
        story.append(Spacer(1, 6))
        story.append(Paragraph(escape(line[2:]), styles["H1Custom"]))
    elif line.startswith("## "):
        story.append(Spacer(1, 4))
        story.append(Paragraph(escape(line[3:]), styles["H2Custom"]))
    elif line.startswith("### "):
        story.append(Paragraph(escape(line[4:]), styles["H2Custom"]))
    elif line.startswith("|"):
        story.append(Preformatted(line, styles["CodeCustom"]))
    elif not line.strip():
        story.append(Spacer(1, 5))
    else:
        story.append(Paragraph(escape(line), styles["BodyCustom"]))

pdf = SimpleDocTemplate(str(PDF_OUT), pagesize=A4, title="Project Documentation", author="GitHub Copilot")
pdf.build(story)

print(f"Wrote {DOCX_OUT}")
print(f"Wrote {PDF_OUT}")
