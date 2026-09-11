from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

root = Path(r'c:/Users/gorla/GitHub Projects/RAG_Assistant')
md_path = root / 'Project_Documentation.md'
out_path = root / 'Project_Documentation.docx'

lines = md_path.read_text(encoding='utf-8').splitlines()

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Project Documentation').italic = True

for line in lines:
    if line.startswith('```'):
        continue
    if line.startswith('# '):
        doc.add_heading(line[2:], level=1)
    elif line.startswith('## '):
        doc.add_heading(line[3:], level=2)
    elif line.startswith('### '):
        doc.add_heading(line[4:], level=3)
    elif not line.strip():
        doc.add_paragraph('')
    else:
        doc.add_paragraph(line)

doc.save(out_path)
print(out_path)
